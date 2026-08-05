import io
import random

from openpyxl import load_workbook

import local_fake
from local_fake import LocalFake


def _draw_sequence(local: LocalFake) -> list:
    return [
        local.ci.first_name(),
        local.ci.last_name(),
        local.ci.phone_number(),
        local.ci.cni_number(),
        local.ci.password(),
        local.ci.address(),
        local.first_name(),  # délégation générique
    ]


def test_same_seed_produces_identical_sequences() -> None:
    local_fake.seed(42)
    first_run = _draw_sequence(LocalFake())

    local_fake.seed(42)
    second_run = _draw_sequence(LocalFake())

    assert first_run == second_run


def test_localfake_constructor_seed_is_equivalent_to_calling_seed() -> None:
    first = LocalFake(seed=7).ci.phone_number()
    second = LocalFake(seed=7).ci.phone_number()
    assert first == second


def test_different_seeds_produce_different_results() -> None:
    local_fake.seed(1)
    first = LocalFake().ci.password()

    local_fake.seed(2)
    second = LocalFake().ci.password()

    assert first != second


def test_seed_does_not_affect_the_global_random_module() -> None:
    random.seed(1234)
    state_before = random.getstate()

    local_fake.seed(555)
    LocalFake().ci.first_name()

    assert random.getstate() == state_before


def test_svg_pdf_and_word_bytes_are_reproducible_under_seed() -> None:
    # SVG (texte pur), PDF (fpdf2) et Word (python-docx, sur notre usage sans
    # mise en forme complexe) sont reproductibles à l'octet près.
    local_fake.seed(99)
    local1 = LocalFake()
    svg1 = local1.file.svg()
    pdf1 = local1.file.pdf(min_size=5_000, max_size=8_000)
    word1 = local1.file.word(min_size=40_000, max_size=45_000)

    local_fake.seed(99)
    local2 = LocalFake()
    svg2 = local2.file.svg()
    pdf2 = local2.file.pdf(min_size=5_000, max_size=8_000)
    word2 = local2.file.word(min_size=40_000, max_size=45_000)

    assert (svg1, pdf1, word1) == (svg2, pdf2, word2)


def test_excel_data_content_is_reproducible_under_seed() -> None:
    # openpyxl a un ordre de sérialisation interne (probablement lié au
    # hachage par identité d'objets de style) qui peut varier à l'octet près
    # selon ce qui a été alloué plus tôt dans le même process, indépendamment
    # de la graine — voir la note dans le README. La donnée elle-même (ce qui
    # compte pour des assertions de test) reste, elle, garantie reproductible.
    local_fake.seed(99)
    local1 = LocalFake()
    excel1 = local1.file.excel(data={"Nom": [local1.ci.last_name() for _ in range(5)]})

    local_fake.seed(99)
    local2 = LocalFake()
    excel2 = local2.file.excel(data={"Nom": [local2.ci.last_name() for _ in range(5)]})

    rows1 = list(load_workbook(io.BytesIO(excel1))["Data"].iter_rows(values_only=True))
    rows2 = list(load_workbook(io.BytesIO(excel2))["Data"].iter_rows(values_only=True))
    assert rows1 == rows2


def test_uniqueness_still_works_after_seeding() -> None:
    local_fake.seed(2026)
    local = LocalFake()
    seen = {local.unique.ci.first_name() for _ in range(8)}
    assert len(seen) == 8
