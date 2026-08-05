import io
import sys

import pytest
from openpyxl import load_workbook

from local_fake import LocalFake
from local_fake.exceptions import InvalidFileSizeError, MissingOptionalDependencyError


@pytest.fixture
def local() -> LocalFake:
    return LocalFake()


def test_svg_respects_custom_size_range(local: LocalFake) -> None:
    content = local.file.svg(min_size=500, max_size=800)
    assert 500 <= len(content) <= 800
    assert content.startswith(b"<svg")


def test_svg_default_size_is_within_default_range(local: LocalFake) -> None:
    content = local.file.svg()
    assert 1_000 <= len(content) <= 3_000


def test_svg_raises_when_max_size_is_unreachably_small(local: LocalFake) -> None:
    with pytest.raises(InvalidFileSizeError):
        local.file.svg(min_size=10, max_size=50)


def test_png_default_size_is_within_default_range(local: LocalFake) -> None:
    content = local.file.png()
    assert 20_000 <= len(content) <= 80_000
    assert content.startswith(b"\x89PNG")


def test_png_with_explicit_dimensions_matches_requested_range(local: LocalFake) -> None:
    content = local.file.png(width=32, height=32, min_size=1_000, max_size=5_000)
    assert 1_000 <= len(content) <= 5_000


def test_png_raises_when_explicit_dimensions_fall_outside_requested_range(local: LocalFake) -> None:
    with pytest.raises(InvalidFileSizeError):
        local.file.png(width=32, height=32, min_size=20_000, max_size=80_000)


def test_pdf_respects_custom_size_range(local: LocalFake) -> None:
    content = local.file.pdf(min_size=5_000, max_size=8_000)
    assert 5_000 <= len(content) <= 8_000
    assert content.startswith(b"%PDF")


def test_word_respects_custom_size_range(local: LocalFake) -> None:
    content = local.file.word(min_size=40_000, max_size=45_000)
    assert 40_000 <= len(content) <= 45_000
    assert content.startswith(b"PK")  # docx = zip


def test_word_raises_when_max_size_is_below_the_minimal_template_size(local: LocalFake) -> None:
    with pytest.raises(InvalidFileSizeError):
        local.file.word(min_size=20_000, max_size=25_000)


def test_excel_with_no_data_returns_empty_sheet(local: LocalFake) -> None:
    content = local.file.excel(min_size=6_000, max_size=8_000)
    assert 6_000 <= len(content) <= 8_000
    assert content.startswith(b"PK")  # xlsx = zip


def test_excel_with_list_data_uses_it_as_header_row(local: LocalFake) -> None:
    content = local.file.excel(data=["Nom", "Prénom", "Téléphone"])
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Data"]
    assert [cell.value for cell in sheet[1]] == ["Nom", "Prénom", "Téléphone"]
    assert sheet.max_row == 1


def test_excel_with_dict_data_uses_keys_as_headers_and_values_as_rows(local: LocalFake) -> None:
    content = local.file.excel(data={"Nom": ["Koné", "Bamba"], "Age": [25, 31]})
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Data"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Nom", "Age")
    assert rows[1] == ("Koné", 25)
    assert rows[2] == ("Bamba", 31)


def test_excel_with_dict_data_pads_shorter_columns_with_blank_cells(local: LocalFake) -> None:
    content = local.file.excel(data={"A": [1, 2, 3], "B": [1]})
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Data"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows == [("A", "B"), (1, 1), (2, None), (3, None)]


def test_excel_rejects_invalid_data_type(local: LocalFake) -> None:
    with pytest.raises(TypeError):
        local.file.excel(data="not a list or dict")


def test_file_path_argument_writes_to_disk(local: LocalFake, tmp_path) -> None:
    target = tmp_path / "generated.svg"
    content = local.file.svg(path=target)
    assert target.exists()
    assert target.read_bytes() == content


def test_min_size_greater_than_max_size_raises() -> None:
    local = LocalFake()
    with pytest.raises(ValueError):
        local.file.svg(min_size=5_000, max_size=1_000)


@pytest.mark.parametrize(
    ("method_name", "kwargs", "missing_module"),
    [
        ("excel", {}, "openpyxl"),
        ("word", {}, "docx"),
        ("pdf", {}, "fpdf"),
        ("png", {}, "PIL.Image"),
    ],
)
def test_missing_optional_dependency_raises_a_clear_error(
    local: LocalFake, monkeypatch: pytest.MonkeyPatch, method_name: str, kwargs: dict, missing_module: str
) -> None:
    monkeypatch.setitem(sys.modules, missing_module, None)
    with pytest.raises(MissingOptionalDependencyError, match=r"pip install 'local-fake\[files\]'"):
        getattr(local.file, method_name)(**kwargs)


def test_svg_does_not_require_any_optional_dependency(local: LocalFake, monkeypatch: pytest.MonkeyPatch) -> None:
    for missing_module in ("openpyxl", "docx", "fpdf", "PIL", "PIL.Image"):
        monkeypatch.setitem(sys.modules, missing_module, None)
    content = local.file.svg()
    assert content.startswith(b"<svg")
